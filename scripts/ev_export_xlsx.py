#!/usr/bin/env python3
"""Export /config/ev_sessions.csv -> naformátovaný /config/www/ev_sessions.xlsx.

Volané z shell_command.ev_export_xlsx a z ev_log_session.sh (po každej session).
openpyxl sa pri chýbaní doinštaluje (HAOS python nemá xlsx knižnice natívne).
"""
import csv
import os

CSV = "/config/ev_sessions.csv"
OUT = "/config/www/ev_sessions.xlsx"

try:
    import openpyxl
except ImportError:
    os.system("pip3 install openpyxl --break-system-packages -q")
    import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

NUMERIC = {"energy_kwh", "charging_min", "cost_eur", "avg_price", "solar_kwh",
           "grid_kwh", "solar_pct", "t1_kwh", "t2_kwh", "t3_kwh", "t4_kwh",
           "soc_start", "soc_end", "odo_start", "odo_end"}

HEADERS_SK = {
    "session_id": "ID", "start": "Začiatok", "end": "Koniec", "user": "Auto",
    "energy_kwh": "kWh", "charging_min": "Min", "cost_eur": "€", "avg_price": "€/kWh",
    "solar_kwh": "Solar kWh", "grid_kwh": "Sieť kWh", "solar_pct": "Solar %",
    "t1_kwh": "T1", "t2_kwh": "T2", "t3_kwh": "T3", "t4_kwh": "T4",
    "soc_start": "SOC od", "soc_end": "SOC do", "odo_start": "km od",
    "odo_end": "km do", "stop_reason": "Dôvod konca",
}


def num(v):
    try:
        f = float(v)
        return int(f) if f.is_integer() else round(f, 3)
    except (ValueError, TypeError):
        return None if v in (None, "") else v


def main():
    if not os.path.exists(CSV):
        print("no csv")
        return
    with open(CSV, newline="", encoding="utf-8") as fh:
        rd = csv.DictReader(fh)
        fieldnames = rd.fieldnames or []
        rows = [r for r in rd if r.get("session_id")]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sessions"
    hfont = Font(bold=True, color="FFFFFF")
    hfill = PatternFill("solid", fgColor="2563EB")
    center = Alignment(horizontal="center")

    for ci, key in enumerate(fieldnames, 1):
        c = ws.cell(1, ci, HEADERS_SK.get(key, key))
        c.font = hfont
        c.fill = hfill
        c.alignment = center

    for ri, r in enumerate(rows, 2):
        for ci, key in enumerate(fieldnames, 1):
            v = r.get(key, "")
            ws.cell(ri, ci, num(v) if key in NUMERIC else v)

    for ci, key in enumerate(fieldnames, 1):
        w = max(len(str(HEADERS_SK.get(key, key))), 9)
        if key in ("start", "end"):
            w = 18
        elif key == "session_id":
            w = 16
        ws.column_dimensions[get_column_letter(ci)].width = w + 2

    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(fieldnames))}{len(rows) + 1}"

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    print(f"wrote {OUT} ({len(rows)} rows)")


if __name__ == "__main__":
    main()
